"""
ETL - Integração Censo Escolar 2025 + Ideb 2025 + INSE 2023
-------------------------------------------------------------
TCC: Visualização Interativa de Dados Educacionais

Escopo: somente escolas da rede PÚBLICA (Federal, Estadual, Municipal).
Escolas privadas são excluídas desde a leitura do Censo Escolar, pois não
possuem Ideb nem INSE divulgados nas bases oficiais utilizadas.

Lê:
  - Tabela_Escola_2025_V2.csv      (Censo Escolar - dados da escola)
  - Tabela_Matricula_2025_V2.csv   (Censo Escolar - matrículas)
  - divulgacao_anos_iniciais_escolas_2025.xlsx  (Ideb)
  - divulgacao_anos_finais_escolas_2025.xlsx    (Ideb)
  - divulgacao_ensino_medio_escolas_2025.xlsx   (Ideb)
  - INSE_2023_escolas_INSE_ESC_2023.csv         (INSE - nível socioeconômico)

Gera:
  - dados_tratados/escolas_ideb_infra_2025.csv
    (uma linha por escola x etapa de ensino, já integrada e tipada)

Ajuste as constantes RAW_DIR / OUT_DIR abaixo conforme sua pasta local.
"""

import pandas as pd
import openpyxl
from pathlib import Path

# ----------------------------------------------------------------------
# Configuração de pastas
# ----------------------------------------------------------------------
RAW_DIR = Path("dados_brutos")       # onde estão os CSVs e as pastas extraídas dos .zip do Ideb
OUT_DIR = Path("dados_tratados")
OUT_DIR.mkdir(exist_ok=True)

IDEB_FILES = {
    "anos_iniciais": RAW_DIR / "divulgacao_anos_iniciais_escolas_2025" / "divulgacao_anos_iniciais_escolas_2025.xlsx",
    "anos_finais":   RAW_DIR / "divulgacao_anos_finais_escolas_2025" / "divulgacao_anos_finais_escolas_2025.xlsx",
    "ensino_medio":  RAW_DIR / "divulgacao_ensino_medio_escolas_2025" / "divulgacao_ensino_medio_escolas_2025.xlsx",
}
ANO_IDEB = "2025"

CSV_ESCOLA = RAW_DIR / "Tabela_Escola_2025_V2.csv"
CSV_MATRICULA = RAW_DIR / "Tabela_Matricula_2025_V2.csv"
CSV_INSE = RAW_DIR / "INSE_2023_escolas_INSE_ESC_2023.csv"  # ajuste o nome se o seu arquivo baixado vier diferente


# ----------------------------------------------------------------------
# 1. IDEB (um arquivo xlsx por etapa de ensino)
# ----------------------------------------------------------------------
# As planilhas de divulgação do Ideb têm um cabeçalho em múltiplas linhas
# (título mesclado + linha técnica). A linha técnica é identificada pela
# primeira célula ser "SG_UF".
def _find_header_row(ws, key_col="SG_UF", max_scan=15):
    for i, row in enumerate(ws.iter_rows(values_only=True, max_row=max_scan)):
        if row and row[0] == key_col:
            return i
    raise ValueError(f"Linha de cabeçalho (coluna '{key_col}') não encontrada nas primeiras {max_scan} linhas.")


def load_ideb_etapa(path: Path, etapa: str, ano: str = ANO_IDEB) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_row_idx = _find_header_row(ws)

    header = None
    data_rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == header_row_idx:
            header = row
        elif i > header_row_idx:
            data_rows.append(row)

    df = pd.DataFrame(data_rows, columns=header)

    # Colunas de interesse. VL_OBSERVADO_<ano> é o valor do Ideb propriamente dito;
    # VL_NOTA_MEDIA e VL_INDICADOR_REND são os dois componentes que formam o Ideb
    # (proficiência no Saeb e taxa de rendimento/aprovação).
    keep = {
        "ID_ESCOLA": "co_entidade",
        "REDE": "rede_ideb",
        f"VL_OBSERVADO_{ano}": "ideb",
        f"VL_NOTA_MEDIA_{ano}": "saeb_nota_media",
        f"VL_INDICADOR_REND_{ano}": "indicador_rendimento",
    }
    df = df[list(keep.keys())].rename(columns=keep)
    df["etapa_ensino"] = etapa

    # Valores ausentes/sigilosos vêm como "-", "*" etc. -> viram NaN com to_numeric
    df["co_entidade"] = pd.to_numeric(df["co_entidade"], errors="coerce")
    for c in ["ideb", "saeb_nota_media", "indicador_rendimento"]:
        df[c] = pd.to_numeric(df[c], errors="coerce")

    df = df.dropna(subset=["co_entidade"])
    df["co_entidade"] = df["co_entidade"].astype("int64")
    return df


def load_ideb() -> pd.DataFrame:
    print("Lendo bases do Ideb...")
    partes = [load_ideb_etapa(path, etapa) for etapa, path in IDEB_FILES.items()]
    ideb = pd.concat(partes, ignore_index=True)
    print(f"  {len(ideb):,} registros (escola x etapa)")
    return ideb


# ----------------------------------------------------------------------
# 2. Censo Escolar - dados da escola (infraestrutura/tecnologia)
# ----------------------------------------------------------------------
COLS_ESCOLA = [
    "CO_ENTIDADE", "NO_ENTIDADE", "SG_UF", "CO_MUNICIPIO", "NO_MUNICIPIO",
    "TP_DEPENDENCIA", "TP_LOCALIZACAO", "TP_SITUACAO_FUNCIONAMENTO",
    # --- infraestrutura tecnológica ---
    "IN_INTERNET", "IN_BANDA_LARGA", "IN_LABORATORIO_INFORMATICA",
    "IN_LABORATORIO_CIENCIAS", "IN_COMPUTADOR",
    "QT_DESKTOP_ALUNO", "QT_TABLET_ALUNO",
    "IN_INTERNET_ALUNOS", "IN_ACESSO_INTERNET_COMPUTADOR",
    # --- infraestrutura básica (saneamento) ---
    "IN_AGUA_POTAVEL", "IN_ENERGIA_REDE_PUBLICA", "IN_ESGOTO_REDE_PUBLICA",
    # --- infraestrutura física / pedagógica não tecnológica ---
    "IN_BANHEIRO", "IN_BANHEIRO_PNE", "IN_BIBLIOTECA", "IN_BIBLIOTECA_SALA_LEITURA",
    "IN_QUADRA_ESPORTES", "IN_COZINHA", "IN_REFEITORIO", "IN_PATIO_COBERTO",
    "IN_PARQUE_INFANTIL", "IN_SALA_ATENDIMENTO_ESPECIAL", "IN_ALIMENTACAO",
    # --- acessibilidade (agregada em uma coluna derivada, ver abaixo) ---
    "IN_ACESSIBILIDADE_RAMPAS", "IN_ACESSIBILIDADE_CORRIMAO",
    "IN_ACESSIBILIDADE_ELEVADOR", "IN_ACESSIBILIDADE_PISOS_TATEIS",
    "IN_ACESSIBILIDADE_SINAL_SONORO", "IN_ACESSIBILIDADE_SINAL_TATIL",
    "IN_ACESSIBILIDADE_SINAL_VISUAL", "IN_ACESSIBILIDADE_INEXISTENTE",
]

# Colunas de acessibilidade que, se qualquer uma for 1, indicam que a escola
# possui ao menos um recurso de acessibilidade (evita 8 colunas binárias soltas).
COLS_ACESSIBILIDADE = [
    "in_acessibilidade_rampas", "in_acessibilidade_corrimao",
    "in_acessibilidade_elevador", "in_acessibilidade_pisos_tateis",
    "in_acessibilidade_sinal_sonoro", "in_acessibilidade_sinal_tatil",
    "in_acessibilidade_sinal_visual",
]

MAPA_DEPENDENCIA = {1: "Federal", 2: "Estadual", 3: "Municipal", 4: "Privada"}
MAPA_LOCALIZACAO = {1: "Urbana", 2: "Rural"}


def load_escolas() -> pd.DataFrame:
    print("Lendo Tabela_Escola...")
    escolas = pd.read_csv(CSV_ESCOLA, sep=";", encoding="latin1", usecols=COLS_ESCOLA)
    escolas.columns = escolas.columns.str.lower()

    # manter apenas escolas em atividade (TP_SITUACAO_FUNCIONAMENTO == 1)
    escolas = escolas[escolas["tp_situacao_funcionamento"] == 1].copy()

    # manter apenas rede PÚBLICA (Federal=1, Estadual=2, Municipal=3).
    # Escolas privadas (código 4) são excluídas aqui: elas não têm Ideb nem
    # INSE divulgados nas bases oficiais, então não agregam à análise.
    escolas = escolas[escolas["tp_dependencia"].isin([1, 2, 3])].copy()

    escolas["rede_ensino"] = escolas["tp_dependencia"].map(MAPA_DEPENDENCIA)
    escolas["localizacao"] = escolas["tp_localizacao"].map(MAPA_LOCALIZACAO)

    # Indicador agregado: possui ao menos um recurso de acessibilidade (1) ou não (0).
    escolas["in_acessibilidade_algum_recurso"] = (
        escolas[COLS_ACESSIBILIDADE].fillna(0).sum(axis=1) > 0
    ).astype(int)

    print(f"  {len(escolas):,} escolas públicas em atividade")
    return escolas


# ----------------------------------------------------------------------
# 3. INSE 2023 - Indicador de Nível Socioeconômico
# ----------------------------------------------------------------------
# Publicado a cada edição do Saeb (anos ímpares). 2023 é a edição mais
# recente disponível — o INSE referente ao Saeb 2025 ainda não havia sido
# divulgado no momento deste trabalho. Une-se pelo mesmo código Inep
# (ID_ESCOLA = CO_ENTIDADE). Cobre apenas escolas públicas, o que já é
# compatível com o filtro aplicado em load_escolas().
COLS_INSE = ["ID_ESCOLA", "MEDIA_INSE", "INSE_CLASSIFICACAO", "QTD_ALUNOS_INSE"]


def load_inse() -> pd.DataFrame:
    print("Lendo INSE 2023...")
    inse = pd.read_csv(CSV_INSE, sep=";", encoding="latin1", decimal=",", usecols=COLS_INSE)
    inse = inse.rename(columns={
        "ID_ESCOLA": "co_entidade",
        "MEDIA_INSE": "inse_media",
        "INSE_CLASSIFICACAO": "inse_classificacao",
        "QTD_ALUNOS_INSE": "inse_qtd_alunos",
    })
    # uma escola aparece uma única vez no arquivo do INSE; ainda assim,
    # agrupamos por segurança caso haja duplicidade de linhas
    inse = inse.groupby("co_entidade", as_index=False).first()
    print(f"  {len(inse):,} escolas com INSE calculado")
    return inse


# ----------------------------------------------------------------------
# 4. Censo Escolar - matrícula (para computar indicadores per capita)
# ----------------------------------------------------------------------
COLS_MATRICULA = ["CO_ENTIDADE", "QT_MAT_BAS", "QT_MAT_FUND_AI", "QT_MAT_FUND_AF", "QT_MAT_MED"]


def load_matricula() -> pd.DataFrame:
    print("Lendo Tabela_Matricula...")
    matricula = pd.read_csv(CSV_MATRICULA, sep=";", encoding="latin1", usecols=COLS_MATRICULA)
    matricula.columns = matricula.columns.str.lower()
    # uma escola pode aparecer em mais de uma linha (turno/modalidade distintos) -> soma
    matricula = matricula.groupby("co_entidade", as_index=False).sum(numeric_only=True)
    print(f"  {len(matricula):,} escolas com matrícula")
    return matricula


# ----------------------------------------------------------------------
# 5. Integração e indicadores derivados
# ----------------------------------------------------------------------
def montar_base() -> pd.DataFrame:
    ideb = load_ideb()
    escolas = load_escolas()
    matricula = load_matricula()
    inse = load_inse()

    base = escolas.merge(matricula, on="co_entidade", how="left")
    # left join: nem toda escola pública tem INSE (ex.: sem alunos suficientes
    # no questionário do Saeb 2023), então mantemos a escola mesmo sem o indicador
    base = base.merge(inse, on="co_entidade", how="left")

    base["computadores_por_aluno"] = base["qt_desktop_aluno"] / base["qt_mat_bas"].replace(0, pd.NA)
    base["tablets_por_aluno"] = base["qt_tablet_aluno"] / base["qt_mat_bas"].replace(0, pd.NA)

    final = ideb.merge(base, on="co_entidade", how="inner")
    pct_com_inse = final["inse_media"].notna().mean() * 100
    print(f"Base final: {len(final):,} linhas (escola x etapa), "
          f"{final['co_entidade'].nunique():,} escolas únicas, "
          f"{pct_com_inse:.1f}% com INSE preenchido")
    return final


if __name__ == "__main__":
    final = montar_base()
    out_path = OUT_DIR / "escolas_ideb_infra_2025.csv"
    final.to_csv(out_path, index=False, encoding="utf-8")
    print(f"Arquivo salvo em: {out_path}")
